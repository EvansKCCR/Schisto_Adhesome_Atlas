#!/usr/bin/env python3
"""Architecture-first prioritization adapted from the curator's supplied script.

WSL: python3 prioritize_candidates.py Output.xlsx -o candidate_priorities.xlsx
Atlas: rank_catalogue(candidates, measured_nodes, network_mode='Measured topology').
No argument reads the current Atlas workbooks and derived alignment evidence.
Missing components remain missing; priority_score is their available-evidence
lower bound and priority_score_upper adds the maximum unassessed contribution.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
from pathlib import Path
import re

import pandas as pd

VERSION = '2.0'
SOURCE_SCRIPT_SHA256 = '2624bf39f45b9e5f663713e6209e9deb5f6e6ad7243e0769fd8d1e213fe136e8'
GRADE_THRESHOLDS = {'A': 8., 'B': 5., 'C': 3.}
PRIORITY_THRESHOLDS = {'Priority 1': 16., 'Priority 2': 13., 'Priority 3': 10.}
SPECIES = ('Shae', 'Sjap', 'Sman')
SPECIES_NAMES = {'Shae': 'S. haematobium', 'Sjap': 'S. japonicum', 'Sman': 'S. mansoni'}
ADHESOME_CENTRALITY_PROXY = {
    'integrin_alpha': 2., 'integrin_beta': 2., 'talin': 2., 'kindlin': 2.,
    'PINCH_like': 2., 'paxillin_like': 2., 'FAK': 2., 'Src': 1.5,
    'vinculin': 1.5, 'ILK': 1.5, 'Nck': 1.5, 'alpha_actinin': 1.,
    'filamin': 1., 'laminin_like': 1., 'collagen_like': .5,
}
INTERFACE_RELEVANCE = {
    'integrin_receptors': 2., 'membrane_proximal_adaptors': 2.,
    'focal_adhesion_scaffolds': 2., 'extracellular_matrix': 1.5,
    'actin_linking': 1., 'downstream_signalling': 1.,
}
FN_ARCHITECTURE_RELEVANCE = {
    'single_pass_FN3_adhesion_receptor_like': 2., 'receptor_protein_tyrosine_phosphatase_like': 2.,
    'receptor_kinase_like': 2., 'secreted_FN3_protein_like': 1.5,
    'single_pass_FN3_membrane_protein_like': 1.5, 'multipass_membrane_FN3_protein_like': 1.,
    'protein_kinase_FN3_like': 1., 'C_type_lectin_Ig_FN3_like': 1.,
    'intracellular_FN3_domain_protein_like': .5, 'TRIM_like_FN3_protein': .5, 'TNFRSF_like': .5,
}
BALANCED_PANEL_ROLES = [
    ('integrin_beta', 'receptor'), ('integrin_alpha', 'receptor'),
    ('talin', 'membrane-proximal adaptor'), ('kindlin', 'membrane-proximal adaptor'),
    ('PINCH_like', 'scaffold'), ('paxillin_like', 'scaffold'), ('ILK', 'signalling'),
    ('Nck', 'signalling'), ('FAK', 'signalling'), ('Src', 'signalling'),
]
FN_PRIORITIES = ['Schistosoma-only single-pass FN3 receptors', 'Schistosoma-only RPTP-like proteins', 'Conserved secreted FN3 proteins']
INTEGRIN_ALPHA_GROUP = 'Integrin alpha (OG0000401 + OG0001220)'
ADH_COMPONENTS = {'domain_score': 4., 'orthology_score': 3., 'motif_score': 2., 'localization_score': 1.,
                  'conservation_score': 2., 'interface_score': 2., 'network_score': 2.,
                  'host_divergence_score': 2., 'tractability_score': 2.}
FN_COMPONENTS = {'evidence_score': 10., 'conservation_score': 2., 'fn_architecture_score': 2.,
                 'fn_motif_context_score': 2., 'host_divergence_score': 2., 'tractability_score': 2.}


def text(value):
    return '' if value is None or pd.isna(value) else str(value).strip()


def first(row, *names):
    return next((row[name] for name in names if name in row and text(row[name])), None)


def number(value, default=float('nan')):
    try:
        result = float(value)
        return result if math.isfinite(result) else default
    except (TypeError, ValueError):
        return default


def bounded(value, maximum, name):
    result = number(value)
    if math.isfinite(result) and not 0 <= result <= maximum:
        raise ValueError(f'{name} must be between 0 and {maximum}: {value}')
    if text(value) and not math.isfinite(result):
        raise ValueError(f'{name} must be numeric: {value}')
    return result


def species_from_id(identifier, fallback=''):
    code = text(identifier).split('__')[0]
    if code in SPECIES:
        return code
    return next((key for key, label in SPECIES_NAMES.items() if text(fallback) in (key, label)), 'Unknown')


def conservation(row):
    counts = [number(first(row, f'Orthology_{species}_HOG_copy_count', f'{species}_HOG_copy_counts')) for species in SPECIES]
    present = sum(value > 0 for value in counts) if all(math.isfinite(value) for value in counts) else number(row.get('schistosome_species_in_HOG'))
    if not math.isfinite(present):
        return float('nan'), 'Schistosome HOG copy-count coverage unavailable'
    if not 0 <= present <= 3 or any(value < 0 for value in counts if math.isfinite(value)):
        raise ValueError('Invalid schistosome HOG copy counts')
    return (2. if present == 3 else 1. if present == 2 else 0.), f'{int(present)}/3 schistosome species in source HOG'


def host_selectivity(row):
    aligned = number(row.get('host_divergence'))
    if math.isfinite(aligned):
        bounded(aligned, 1., 'host_divergence')
        return 2 * aligned, 'Aligned human-homologue divergence × 2; ' + text(row.get('host_divergence_basis'))
    direct = text(first(row, 'Orthology_Hsap_direct_orthologues', 'Hsap_orthologues'))
    relationship = text(first(row, 'Orthology_Hsap_relationship', 'Hsap_relationship'))
    count = number(first(row, 'Orthology_Hsap_HOG_copy_count', 'Hsap_HOG_copy_counts'))
    mapping = text(first(row, 'Orthology_mapping_status', 'orthology_mapping_status')).lower()
    negative = re.search(r'no direct|not recovered|not detected|none detected', (direct + ' ' + relationship).lower())
    if negative or (mapping == 'matched' and count == 0 and not direct):
        return 2., 'Orthology selectivity proxy: no direct human orthologue recovered in sampled references'
    if direct and direct.lower() not in ('none', '0', 'nan'):
        return (1. if 'many' in relationship.lower() else 0.), 'Orthology selectivity proxy: ' + (relationship or 'direct human orthologue reported')
    return float('nan'), 'Human alignment and direct-orthologue assessment unavailable'


def tractability(row):
    architecture = text(first(row, 'architecture_evidence', 'architecture', 'NCBI_CDD_architecture'))
    context = ' '.join(text(row.get(field)) for field in ['status', 'length_check', 'rationale']) + ' ' + architecture
    if re.search(r'incomplete|fragment|truncat', context.lower()):
        return 0., 'Recorded incomplete/fragmented model'
    length = number(row.get('length'))
    if not math.isfinite(length):
        return float('nan'), 'Protein length unavailable'
    return (2. if 100 <= length <= 1800 else 1. if 50 <= length <= 4000 else 0.), f'Length-based assay proxy: {int(length)} aa'


def conflict_penalty(row):
    # Do not penalize a resolved competing assignment or a negated conflict.
    context = ' '.join(text(row.get(field)) for field in ['rationale', 'architecture_evidence', 'motif_context_summary', 'MotifPrediction_context_summary', 'topology_evidence', 'status'])
    context = re.sub(r'(?:no|without) (?:evidence of )?(?:localization|topology|context) conflict', '', context.lower())
    reasons = []
    if re.search(r'localization conflict|context_conflict|context is conflicting|topology conflict|discordance', context):
        reasons.append('Context/topology conflict')
    if re.search(r'incomplete|fragment|truncat', context):
        reasons.append('Incomplete protein model')
    if re.search(r'unresolved paralog|ambiguous/manual|unresolved competing', context):
        reasons.append('Unresolved paralogue/family assignment')
    if re.search(r'architecture incompatible|incompatible architecture|defining architecture absent', context):
        return 3., 'Incompatible defining architecture'
    return min(3., float(len(reasons))), '; '.join(reasons) or 'No recorded conflict'


def resolve_grade(row, evidence, domain, minimum):
    decision = text(row.get('family_decision'))
    context = decision + ' ' + text(row.get('architecture_evidence')) + ' ' + text(row.get('rationale'))
    incompatible = bool(re.search(r'do not assign|architecture incompatible|incompatible architecture|defining architecture absent', context.lower()))
    grade = text(row.get('grade')).upper()
    basis = 'Workbook grade with architecture gate'
    if grade not in 'ABCD' or len(grade) != 1:
        grade = next((letter for letter, threshold in GRADE_THRESHOLDS.items() if evidence >= threshold), 'D')
        basis = 'Calculated from available base evidence with architecture gate'
    if incompatible:
        grade = 'D'
    elif not math.isfinite(domain) or domain < minimum:
        grade = 'C' if grade in ('A', 'B') else grade
    if decision == 'Family-level provisional only':
        grade = 'C' if grade != 'D' else grade
    hard_fail = incompatible or grade == 'D' or not math.isfinite(domain) or domain < minimum
    return grade, basis, hard_fail


def finish_score(result, row, components, minimum):
    available = {name: result[name] for name in components if math.isfinite(result[name])}
    missing = [name for name in components if name not in available]
    score = sum(available.values()) - result['conflict_penalty']
    upper = score + sum(components[name] for name in missing)
    grade, basis, hard_fail = resolve_grade(row, result['evidence_score'], result['domain_score'], minimum)
    result.update(grade=grade, grade_basis=basis, source_grade=text(row.get('grade')), hard_fail=hard_fail,
                  priority_score=round(score, 3), priority_score_upper=round(upper, 3),
                  score_complete=not missing, missing_components='; '.join(missing),
                  available_components=len(available), total_components=len(components))
    if hard_fail or grade not in ('A', 'B'):
        tier = 'Validation first'
    else:
        tier = next((label for label, threshold in PRIORITY_THRESHOLDS.items() if score >= threshold), 'Evidence needed' if missing and upper >= 10 else 'Hold')
    result['priority_tier'] = tier
    result['eligibility_reason'] = ('Architecture/assignment gate not met' if hard_fail else 'Provisional family candidate: validate assignment first' if grade not in ('A', 'B') else 'Grade A/B with defining architecture retained')
    return result


def common(row, candidate_class):
    identifier = text(first(row, 'sequence_id', 'protein_id'))
    family = text(first(row, 'reviewed_family', 'audit_family', 'recommended_family', 'assigned_family', 'family'))
    group = text(first(row, 'orthogroup', 'Orthology_Orthogroup', 'Orthogroup'))
    code = species_from_id(identifier, first(row, 'standardized_id', 'species', 'Specie'))
    result = dict(candidate_class=candidate_class, sequence_id=identifier, protein_id=identifier,
                  species=code, species_name=SPECIES_NAMES.get(code, code), family=family,
                  module=text(row.get('module')), orthogroup=group, family_orthogroup=INTEGRIN_ALPHA_GROUP if group in {'OG0000401', 'OG0001220'} else group,
                  source=text(row.get('source')), family_decision=text(row.get('family_decision')),
                  audit_classification=text(row.get('audit_classification')) or ('Supported' if text(row.get('grade')) in ('A', 'B') else 'Provisional' if text(row.get('grade')) == 'C' else 'Unassigned'),
                  assignment_interpretation=text(row.get('assignment_interpretation')),
                  candidate_annotation=text(row.get('candidate_annotation')), rationale=text(row.get('rationale')),
                  priority_group=text(row.get('priority_group')))
    for field in ['orthogroup_identity','orthogroup_evolution','reference_binding_partners','putative_function','annotation_confidence','annotation_references','orthogroup_annotation_source']:
        result[field]=text(row.get(field))
    if result['family_decision'] == 'Retain provisionally':
        result['audit_classification'] = 'Supported'
        result['assignment_interpretation'] = 'Parasite-specific supported cluster'
    elif result['family_decision'] == 'Family-level provisional only':
        result['audit_classification'] = 'Provisional'
        result['assignment_interpretation'] = 'Provisional candidate'
    elif result['family_decision'].lower().startswith('do not assign'):
        result['audit_classification'] = 'Unassigned'
    for field in ['protein_annotation_id', 'identifier_mapping_status', 'alternative_protein_ids', 'architecture', 'architecture_evidence', 'domain_architecture', 'domain_architecture_basis', 'motif_conservation', 'motif_conservation_basis', 'motif_context_evidence', 'motif_assignment_tiers', 'motif_hit_details', 'topology_compatibility', 'topology_compatibility_basis', 'phylogenetic_support', 'phylogenetic_support_basis', 'host_divergence', 'host_divergence_basis', 'host_orthology_flag', 'host_similarity_flag', 'interface_connections', 'DeepTMHMM', 'DeepLoc_2.1', 'SignalP.v6.0', 'length', 'identifier', 'degree', 'betweenness', 'closeness', 'network_mapping_basis']:
        result[field] = row.get(field, float('nan') if field in ['domain_architecture', 'motif_conservation', 'topology_compatibility', 'phylogenetic_support', 'host_divergence', 'degree', 'betweenness', 'closeness', 'length'] else '')
    result['conservation_score'], result['conservation_basis'] = conservation(row)
    result['host_divergence_score'], result['host_selectivity_basis'] = host_selectivity(row)
    result['tractability_score'], result['tractability_basis'] = tractability(row)
    result['conflict_penalty'], result['conflict_basis'] = conflict_penalty(row)
    return result


def score_adhesome(row, network_mode='Measured topology'):
    result = common(row, 'adhesome')
    for target, source, maximum in [('domain_score', 'domain_score_0_4', 4.), ('orthology_score', 'orthology_score_0_3', 3.), ('motif_score', 'motif_score_0_2', 2.), ('localization_score', 'localization_score_0_1', 1.)]:
        result[target] = bounded(row.get(source), maximum, source)
    result['evidence_score'] = sum(value for value in [result[name] for name in ['domain_score', 'orthology_score', 'motif_score', 'localization_score']] if math.isfinite(value))
    result['base_evidence_basis'] = 'Audited domain 0–4 + orthology 0–3 + motif 0–2 + localization 0–1'
    result['interface_score'] = INTERFACE_RELEVANCE.get(result['module'], .5)
    if network_mode == 'Family-role proxy':
        result['network_score'] = ADHESOME_CENTRALITY_PROXY.get(result['family'], .5)
        result['network_score_basis'] = 'Curator family-role proxy; not measured centrality'
    elif network_mode == 'Measured topology':
        result['network_score'] = bounded(row.get('network_score'), 2., 'network_score')
        result['network_score_basis'] = text(row.get('network_score_basis')) or 'No exact candidate mapping in displayed graph'
    else:
        raise ValueError('Unknown network scoring mode')
    return finish_score(result, row, ADH_COMPONENTS, 2.)


def fn_motif_context(row):
    tiers = text(first(row, 'motif_candidate_tiers', 'motif_assignment_tiers'))
    context = text(row.get('motif_context_summary'))
    motifs = text(row.get('motif_annotation_motifs'))
    if re.search(r'context_conflict|review_only', tiers.lower()) or re.search(r'context_conflict|contradicted|context is conflicting', context.lower()):
        return 0., 'Ligand motif context conflicts with assigned topology/region'
    if re.search(r'\b(?:RGD|LDV)\b', motifs.upper()):
        if 'context_supported' in (tiers + context).lower() or 'localization=supported' in context.lower():
            return 2., 'RGD/LDV ligand motif with supported context'
        return 1., 'RGD/LDV ligand motif; sequence-level context'
    value = bounded(row.get('motif_score_0_1'), 1., 'motif_score_0_1')
    return value, 'Workbook FN3 motif score; no RGD/LDV ligand bonus'


def score_fibronectin(row):
    result = common(row, 'fibronectin-like')
    counts = [bounded(row.get(f'FN{n}_count'), 10000., f'FN{n}_count') for n in [1, 2, 3]]
    result.update(dict(zip(['fn1_count', 'fn2_count', 'fn3_count'], counts)))
    result['domain_score'] = min(4., sum(counts)) if all(math.isfinite(value) for value in counts) else float('nan')
    result['evidence_score'] = bounded(row.get('total_score_0_10'), 10., 'total_score_0_10')
    result['base_evidence_basis'] = 'Curated FN3 total_score_0_10; ligand-context bonus scored separately'
    result['fn_architecture_score'] = FN_ARCHITECTURE_RELEVANCE.get(result['family'], .5)
    result['fn_motif_context_score'], result['fn_motif_context_basis'] = fn_motif_context(row)
    if not result['priority_group'] or result['priority_group'] == 'Not in FN3 review collection':
        scope = text(first(row, 'orthology_scope', 'HOG_status'))
        if scope == 'Schistosoma-only in sampled HOG' and result['family'] == 'single_pass_FN3_adhesion_receptor_like':
            result['priority_group'] = FN_PRIORITIES[0]
        elif scope == 'Schistosoma-only in sampled HOG' and result['family'] == 'receptor_protein_tyrosine_phosphatase_like':
            result['priority_group'] = FN_PRIORITIES[1]
        elif scope == 'shared HOG' and result['family'] == 'secreted_FN3_protein_like' and result['conservation_score'] == 2:
            result['priority_group'] = FN_PRIORITIES[2]
        else:
            result['priority_group'] = 'Other / unresolved FN3 candidates'
    return finish_score(result, row, FN_COMPONENTS, 1.)


def attach_network(candidates, measured_nodes):
    """Exact candidate joins only; duplicate/shared representatives remain explicit."""
    rows = []
    if measured_nodes is None or measured_nodes.empty:
        measured_nodes = pd.DataFrame(columns=['identifier', 'sequence_id', 'degree', 'betweenness', 'closeness'])
    maxima = {field: pd.to_numeric(measured_nodes.get(field, pd.Series(dtype=float)), errors='coerce').max() for field in ['degree', 'betweenness', 'closeness']}
    matches = {}
    for record in measured_nodes.to_dict('records'):
        ids = [identifier.strip() for identifier in re.split(r'\s*[|;]\s*', text(record.get('sequence_id'))) if identifier.strip()]
        for identifier in ids:
            matches.setdefault(identifier, []).append(record)
    for candidate in candidates.to_dict('records'):
        row = dict(candidate)
        identifier = text(first(row, 'sequence_id', 'protein_id'))
        mapped = matches.get(identifier, [])
        row['network_score'] = float('nan')
        row['network_score_basis'] = 'No exact candidate mapping in displayed graph'
        row['identifier'] = ' | '.join(sorted({text(record.get('identifier')) for record in mapped}))
        row['network_mapping_basis'] = 'Exact candidate ID → ' + row['identifier'] if mapped else 'Candidate not mapped in displayed graph'
        for field in ['degree', 'betweenness', 'closeness']:
            values = [number(record.get(field)) for record in mapped]
            row[field] = min(values) if values and all(math.isfinite(value) for value in values) else float('nan')
        flags = sorted({text(record.get('host_similarity_flag')) for record in mapped if text(record.get('host_similarity_flag'))})
        if flags:
            row['host_similarity_flag'] = ' | '.join(flags)
        interfaces = [number(record.get('interface_connections')) for record in mapped]
        if interfaces and all(math.isfinite(value) for value in interfaces):
            row['interface_connections'] = min(interfaces)
        if all(math.isfinite(row[field]) for field in maxima):
            active = [field for field, maximum in maxima.items() if math.isfinite(maximum) and maximum > 0]
            row['network_score'] = 2 * sum(row[field] / maxima[field] for field in active) / len(active) if active else 0.
            row['network_score_basis'] = '2 × mean(max-normalized degree, betweenness, closeness); zero-valued graph-wide metrics omitted' + ('; minimum across multiple exact STRING representatives' if len(mapped) > 1 else '')
        for field in ['domain_architecture', 'topology_compatibility', 'host_divergence', 'phylogenetic_support', 'motif_conservation']:
            overrides = [record for record in mapped if text(record.get(field + '_basis')).startswith('Curated override:')]
            if overrides:
                best = min(overrides, key=lambda record: number(record.get(field), default=1.))
                row[field], row[field + '_basis'] = best[field], best[field + '_basis']
        rows.append(row)
    return pd.DataFrame(rows)


def rank(rows):
    if not rows:
        return pd.DataFrame(columns=['rank', 'candidate_class', 'sequence_id', 'species', 'family', 'orthogroup', 'family_orthogroup', 'priority_score', 'priority_tier', 'grade', 'hard_fail', 'priority_group'])
    frame = pd.DataFrame(rows)
    frame = frame.sort_values(['priority_score', 'domain_score', 'conflict_penalty', 'conservation_score', 'phylogenetic_support', 'sequence_id'], ascending=[False, False, True, False, False, True], na_position='last').reset_index(drop=True)
    frame.insert(0, 'rank', range(1, len(frame) + 1))
    return frame


def summarize_orthogroups(ranked):
    rows = []
    if ranked.empty:
        return pd.DataFrame(columns=['candidate_class', 'family', 'priority_group', 'family_orthogroup', 'source_orthogroups', 'species_count', 'member_count', 'best_score', 'best_tier', 'best_protein'])
    for (kind, family, priority, group), members in ranked[ranked.orthogroup.ne('')].groupby(['candidate_class', 'family', 'priority_group', 'family_orthogroup'], dropna=False, sort=True):
        eligible = members[members.grade.isin(['A', 'B']) & ~members.hard_fail]
        best = (eligible if len(eligible) else members).sort_values('rank').iloc[0]
        ids = {species: '; '.join(sorted(set(members.loc[members.species.eq(species), 'sequence_id']))) for species in SPECIES}
        rows.append(dict(candidate_class=kind, family=family, priority_group=priority, family_orthogroup=group,
                         source_orthogroups='; '.join(sorted(set(members.orthogroup))), species_count=sum(bool(value) for value in ids.values()),
                         member_count=members.sequence_id.nunique(), best_score=best.priority_score, best_tier=best.priority_tier,
                         best_protein=best.sequence_id, eligible_members=len(eligible), **{species + '_candidates': value for species, value in ids.items()}))
    return pd.DataFrame(rows).sort_values(['best_score', 'species_count', 'family'], ascending=[False, False, True]).reset_index(drop=True) if rows else pd.DataFrame()


def select_balanced_panel(ranked):
    rows = []
    eligible = ranked[ranked.grade.isin(['A', 'B']) & ~ranked.hard_fail] if len(ranked) else ranked
    summaries = summarize_orthogroups(eligible)
    if summaries.empty:
        summaries = pd.DataFrame(columns=['candidate_class', 'family', 'priority_group'])
    for family, role in BALANCED_PANEL_ROLES:
        choices = summaries[summaries.candidate_class.eq('adhesome') & summaries.family.eq(family)]
        if len(choices):
            best = choices.sort_values(['species_count', 'best_score', 'family_orthogroup'], ascending=[False, False, True]).iloc[0].to_dict()
            rows.append(dict(best, role=role, selection_status='Selected', selection_reason='Highest species coverage, then highest eligible candidate score'))
        else:
            rows.append(dict(candidate_class='adhesome', family=family, role=role, priority_group='', selection_status='No eligible grade A/B candidate', best_protein='', best_tier='Validation first'))
    for group, role in zip(FN_PRIORITIES, ['FN3 surface recognition', 'FN3/phosphatase signalling', 'Secreted FN3 ligand']):
        choices = summaries[summaries.candidate_class.eq('fibronectin-like') & summaries.priority_group.eq(group)]
        if len(choices):
            best = choices.sort_values(['species_count', 'best_score', 'family_orthogroup'], ascending=[False, False, True]).iloc[0].to_dict()
            rows.append(dict(best, role=role, selection_status='Selected', selection_reason='Independent FN3 priority group; species coverage then eligible score'))
        else:
            rows.append(dict(candidate_class='fibronectin-like', priority_group=group, role=role, family='', selection_status='No eligible grade A/B candidate', best_protein='', best_tier='Validation first'))
    panel = pd.DataFrame(rows)
    if len(panel):
        panel.insert(0, 'selection_order', range(1, len(panel) + 1))
    return panel


def parameter_table(network_mode):
    return pd.DataFrame([
        ('Base adhesome evidence', '0–10', 'Domain 0–4 + orthology 0–3 + motif 0–2 + localization 0–1'),
        ('Architecture gate', 'hard rule', 'Adhesome domain ≥2; FN-domain count score ≥1; incompatible/Do not assign → D. Family-level provisional only → C.'),
        ('Base FN3 evidence', '0–10', 'Curated total_score_0_10; explicit FN ligand-context bonus below'),
        ('Species conservation', '0–2', '2 for three schistosomes, 1 for two, 0 for zero/one; source HOG copy counts'),
        ('Interface / FN architecture', '0–2', 'Curator module weights for adhesome; architecture-family relevance for FN3'),
        ('Network mode', '0–2', network_mode + '; measured = 2 × mean of graph-max-normalized degree/betweenness/closeness over nonzero graph-wide metrics'),
        ('Host selectivity', '0–2', '2 × aligned human divergence where available; otherwise explicit direct-orthology proxy (2 no direct recovered, 1 many relationship, 0 direct); blank remains missing'),
        ('Tractability', '0–2', '2 for 100–1800 aa, 1 for 50–4000 aa, otherwise 0; incomplete models 0'),
        ('FN3 ligand-context bonus', '0–2', 'RGD/LDV: 2 supported context, 1 sequence-only, 0 conflict/review-only; other motifs use workbook 0–1 score'),
        ('Conflict penalty', '0–3 subtracted', 'Context/topology, incomplete model, unresolved paralogy; defining incompatible architecture = 3'),
        ('Priority 1 / 2 / 3', '≥16 / ≥13 / ≥10', 'Grade A/B and architecture gate passed; otherwise Validation first'),
        ('Missing components', 'score interval', 'priority_score sums available points minus penalty; upper bound adds maximum missing points; components stay blank. Tier requires its lower bound.'),
        ('Phylogeny and motif conservation', 'supporting evidence', 'Alignment-derived values and bases retained; phylogenetic support used as tie-breaker, not double-counted in base orthology score'),
        ('Panel selection', 'coverage then rank', 'One eligible family group per adhesome role; three independent FN3 priority groups; source OGs retained under alpha-integrin family merge'),
    ], columns=['component', 'range', 'definition'])


def rank_catalogue(candidates, measured_nodes=None, network_mode='Measured topology'):
    prepared = attach_network(candidates, measured_nodes)
    adh, fn = [], []
    for row in prepared.to_dict('records'):
        is_fn = 'fibronectin_like_candidate' in text(row.get('source')) or bool(text(row.get('recommended_family'))) or text(row.get('candidate_class')) == 'fibronectin-like'
        if not text(first(row, 'sequence_id', 'protein_id')):
            continue
        (fn if is_fn else adh).append(score_fibronectin(row) if is_fn else score_adhesome(row, network_mode))
    a, f = rank(adh), rank(fn)
    combined = pd.concat([a, f], ignore_index=True)
    return {'Ranked_adhesome': a, 'Ranked_fibronectin': f, 'Orthogroup_summary': summarize_orthogroups(combined),
            'Experimental_panel': select_balanced_panel(combined), 'Prioritization_parameters': parameter_table(network_mode)}


def workbook_bytes(tables, sources=None, template=None):
    """Portable application export using the Atlas's installed openpyxl dependency.

    The supplied artifact_tool Python package is absent from the WSL/Cloud app.
    Source workbooks are read only; a template is loaded into an output copy.
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = load_workbook(template) if template else Workbook()
    if not template:
        workbook.remove(workbook.active)
    all_tables = dict(tables)
    all_tables.update(sources or {})
    for name, frame in all_tables.items():
        if name in workbook.sheetnames:
            del workbook[name]
        sheet = workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value.item() if hasattr(value, 'item') else value for value in values])
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='064C2C')
            cell.font = Font(bold=True, color='FFFFFF', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 34
        sheet.freeze_panes = 'C2'
        if len(frame.columns):
            sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(frame.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 42 if any(term in column for term in ['basis', 'rationale', 'annotation', 'architecture', 'definition', 'reason', 'missing']) else 25 if any(term in column for term in ['protein', 'sequence', 'orthogroup', 'group']) else max(16, min(28, len(column) + 2))
            for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if isinstance(cell.value, float):
                        cell.number_format = '0.000'
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
        # Empty result tables remain valid header-only sheets.
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', nargs='?', help='Combined input workbook; omit to use the Atlas library')
    parser.add_argument('--fibronectin', help='Separate FN3 workbook, if not in the input workbook')
    parser.add_argument('--network-nodes', help='CSV of displayed topology with exact sequence_id mappings')
    parser.add_argument('--network-mode', choices=['Measured topology', 'Family-role proxy'], default='Measured topology')
    parser.add_argument('-o', '--output', default='candidate_priorities.xlsx')
    args = parser.parse_args()
    output = Path(args.output).resolve()
    inputs = [Path(name).resolve() for name in [args.input, args.fibronectin] if name]
    if output in inputs:
        parser.error('Output must differ from every source workbook.')
    sources = {}
    if args.input:
        books = pd.read_excel(args.input, sheet_name=None)
        sheet = next((name for name in ['adhesome_candidates', 'Audited_family_assignment', 'candidates'] if name in books), None)
        if sheet is None:
            parser.error('Input workbook needs adhesome_candidates or Audited_family_assignment.')
        a = books[sheet].assign(source=Path(args.input).name + ' / ' + sheet, candidate_class='adhesome')
        f = books.get('fibronectin_like', pd.DataFrame()).assign(candidate_class='fibronectin-like')
        if args.fibronectin:
            fn_books = pd.read_excel(args.fibronectin, sheet_name=None)
            fn_sheet = next((name for name in ['Candidates', 'fibronectin_like', 'candidates'] if name in fn_books), None)
            if fn_sheet is None:
                parser.error('FN3 workbook needs Candidates or fibronectin_like.')
            f = fn_books[fn_sheet].assign(source=Path(args.fibronectin).name + ' / ' + fn_sheet, candidate_class='fibronectin-like')
        candidates = pd.concat([a, f], ignore_index=True)
        if len(f):
            sources['Source_fibronectin'] = f
    else:
        try:
            from data import load
        except ModuleNotFoundError as error:
            if error.name != 'data':
                raise
            parser.error('No-input mode requires the Atlas copy in adhesome_atlas. Pass an input workbook for standalone use.')
        cohorts = load()[0]
        a, f = cohorts['Adhesome candidates'], cohorts['FN3 / fibronectin-like review']
        candidates = pd.concat([a, f], ignore_index=True)
        sources = {'Source_adhesome': a, 'Source_fibronectin': f}
    measured = pd.read_csv(args.network_nodes) if args.network_nodes else None
    tables = rank_catalogue(candidates, measured, args.network_mode)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(workbook_bytes(tables, sources, template=args.input))
    print(f'Wrote {output}: {len(tables["Ranked_adhesome"])} adhesome, {len(tables["Ranked_fibronectin"])} FN3 candidates; {len(tables["Experimental_panel"])} panel groups')


if __name__ == '__main__':
    main()
