#!/usr/bin/env python3
"""Conservative, evidence-based audit of adhesome family assignments.

Usage:
  python conservative_adhesome_audit.py INPUT.xlsx OUTPUT.xlsx

The algorithm uses only computational evidence already present in the Candidates
sheet. Experimental validation fields are never used for scoring or assignment.
"""
from __future__ import annotations
import argparse, hashlib, json, sys
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = "1.0.0"
REFERENCE_SPECIES = ("H. sapiens", "C. elegans", "D. melanogaster", "M. musculus", "X. laevis")
REQUIRED_COLUMNS = {
    "Specie", "protein_id", "family", "module", "length_check", "domain_match",
    "domain_type_fraction", "domain_copy_fraction", "Pfam_architecture",
    "NCBI_CDD_architecture", "DeepLoc_2.1", "DeepTMHMM", "SignalP.v6.0",
    "MotifAnnotation_confidence", "MotifAnnotation_candidate_tiers",
    "Orthology_HOG_status", "Orthology_reference_support"
}
AUDIT_COLUMNS = [
    "Audit_domain_score_0_4", "Audit_orthology_score_0_3",
    "Audit_motif_score_0_2", "Audit_localization_score_0_1",
    "Audit_total_score_0_10", "Audit_grade", "Audit_family_decision",
    "Audit_rationale", "Audit_competing_supported_families",
    "Audit_experimental_validation_used"
]

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def number(x, default=0.0):
    try: return float(x)
    except (TypeError, ValueError): return default

def truthy(x):
    if isinstance(x, bool): return x
    if isinstance(x, (int, float)): return x != 0
    return str(x).strip().lower() in {"true", "1", "yes"}

def domain_score(row):
    """Primary gate: complete architecture outranks all supporting evidence."""
    if truthy(row.get("domain_match")): return 4
    fraction = min(number(row.get("domain_type_fraction")), number(row.get("domain_copy_fraction")))
    if fraction >= 0.50: return 2
    if fraction > 0: return 1
    return 0

def orthology_score(row):
    hog = str(row.get("Orthology_HOG_status") or "")
    support = str(row.get("Orthology_reference_support") or "")
    direct_n = sum(species in support for species in REFERENCE_SPECIES)
    if hog == "shared HOG" and direct_n >= 4: return 3
    if hog == "shared HOG" and direct_n >= 1: return 2
    if hog == "shared HOG": return 1
    return 0

def motif_score(row):
    tier = str(row.get("MotifAnnotation_candidate_tiers") or "")
    confidence = str(row.get("MotifAnnotation_confidence") or "")
    if "context_supported_candidate" in tier or confidence in {"Very high", "High"}: return 2
    if "sequence_only_candidate" in tier or confidence == "Medium": return 1
    return 0

def localization_score(row):
    module = str(row.get("module") or "")
    loc = str(row.get("DeepLoc_2.1") or "")
    tm = str(row.get("DeepTMHMM") or "")
    signalp = str(row.get("SignalP.v6.0") or "")
    if module == "extracellular_matrix":
        return int("Extracellular" in loc and ("SP" in tm or signalp.startswith("SP")))
    if module == "integrin_receptors":
        return int("Cell membrane" in loc and "TM" in tm)
    return int(any(label in loc for label in ("Cytoplasm", "Cell membrane", "Nucleus")))

def audit_row(row, alternatives):
    ds, os_, ms, ls = domain_score(row), orthology_score(row), motif_score(row), localization_score(row)
    total = ds + os_ + ms + ls
    complete = sorted({x["family"] for x in alternatives if truthy(x.get("domain_match"))})
    multifamily = len({x["family"] for x in alternatives}) > 1
    reasons = []
    if not truthy(row.get("domain_match")): reasons.append("required family architecture incomplete")
    if row.get("length_check") == "shorter_than_expected": reasons.append("sequence shorter than family threshold")
    if os_ == 0: reasons.append("no informative cross-metazoan orthology support")
    if multifamily: reasons.append("same protein occurs under multiple candidate families")
    if "context_conflict" in str(row.get("MotifAnnotation_candidate_tiers") or ""):
        reasons.append("motif context conflict")
    cdd = str(row.get("NCBI_CDD_architecture") or "")
    family = row["family"]
    if family == "FAK" and not truthy(row.get("domain_match")) and any(x in cdd for x in ("PLK", "AMPK", "SPS1")):
        reasons.append("CDD supports a non-FAK kinase architecture")
    if family == "Src" and not truthy(row.get("domain_match")):
        pfam = str(row.get("Pfam_architecture") or "")
        if not all(x in pfam for x in ("PF00018", "PF00017")):
            reasons.append("Src SH3-SH2-kinase architecture not recovered")
    if family == "ILK" and not truthy(row.get("domain_match")):
        reasons.append("ILK-defining architecture not recovered")

    if multifamily:
        if truthy(row.get("domain_match")) and len(complete) == 1:
            decision = "Retain; resolves competing family"
        elif len(complete) > 1:
            decision = "Ambiguous; manual family resolution"
        elif not truthy(row.get("domain_match")):
            decision = "Do not assign at family level"
        else:
            decision = "Manual review"
    elif truthy(row.get("domain_match")) and os_ >= 2:
        decision = "Retain family assignment"
    elif truthy(row.get("domain_match")):
        decision = "Retain provisionally"
    elif min(number(row.get("domain_type_fraction")), number(row.get("domain_copy_fraction"))) >= 0.5 and os_ >= 2:
        decision = "Family-level provisional only"
    else:
        decision = "Do not assign at family level"

    if decision == "Retain family assignment" and total >= 7: grade = "A"
    elif decision.startswith("Retain") and total >= 5: grade = "B"
    elif decision.startswith("Family-level") or decision.startswith("Ambiguous"): grade = "C"
    else: grade = "D"
    return [ds, os_, ms, ls, total, grade, decision,
            "; ".join(reasons) if reasons else "evidence layers are concordant",
            ", ".join(complete), "No"]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--sheet", default="Candidates")
    ap.add_argument("--manifest", type=Path, default=None)
    args = ap.parse_args()
    if args.input.resolve() == args.output.resolve():
        sys.exit("Refusing to overwrite the input workbook.")
    wb = load_workbook(args.input)
    if args.sheet not in wb.sheetnames: sys.exit(f"Missing sheet: {args.sheet}")
    ws = wb[args.sheet]
    headers = [c.value for c in ws[1]]
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing: sys.exit("Missing required columns: " + ", ".join(missing))
    col = {h: i + 1 for i, h in enumerate(headers)}
    records = []
    for r in range(2, ws.max_row + 1):
        d = {h: ws.cell(r, c).value for h, c in col.items()}; d["_row"] = r; records.append(d)
    by_protein = defaultdict(list)
    for d in records: by_protein[(d["Specie"], d["protein_id"])].append(d)

    start = ws.max_column + 1
    for j, h in enumerate(AUDIT_COLUMNS, start): ws.cell(1, j, h)
    outcomes = []
    for d in records:
        result = audit_row(d, by_protein[(d["Specie"], d["protein_id"])])
        outcomes.append(result)
        for j, value in enumerate(result, start): ws.cell(d["_row"], j, value)

    blue = PatternFill("solid", fgColor="1F4E78")
    for c in range(start, start + len(AUDIT_COLUMNS)):
        ws.cell(1, c).fill = blue; ws.cell(1, c).font = Font(color="FFFFFF", bold=True)
        ws.cell(1, c).alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = 22 if c < start + 6 else 38
    color = {"A":"C6EFCE", "B":"DDEBF7", "C":"FFF2CC", "D":"F4CCCC"}
    for d, result in zip(records, outcomes):
        ws.cell(d["_row"], start + 5).fill = PatternFill("solid", fgColor=color[result[5]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    if "Audit_Summary" in wb.sheetnames: del wb["Audit_Summary"]
    summary = wb.create_sheet("Audit_Summary", 0)
    summary.append(["Metric", "Value"])
    grades = Counter(x[5] for x in outcomes); decisions = Counter(x[6] for x in outcomes)
    metrics = [("Algorithm version", VERSION), ("Candidate rows", len(records)),
               ("Unique proteins", len(by_protein)), ("Input SHA-256", sha256(args.input))]
    metrics += [(f"Grade {g}", grades[g]) for g in "ABCD"]
    metrics += sorted(decisions.items())
    metrics.append(("Experimental validation used", "No"))
    for item in metrics: summary.append(item)
    for cell in summary[1]: cell.fill = blue; cell.font = Font(color="FFFFFF", bold=True)
    summary.column_dimensions["A"].width = 40; summary.column_dimensions["B"].width = 70
    summary.freeze_panes = "A2"; summary.sheet_view.showGridLines = False

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    manifest = args.manifest or args.output.with_suffix(".manifest.json")
    manifest.write_text(json.dumps({
        "algorithm_version": VERSION,
        "command": f"python conservative_adhesome_audit.py {args.input.name} {args.output.name}",
        "input_file": args.input.name, "input_sha256": sha256(args.input),
        "output_file": args.output.name, "output_sha256": sha256(args.output),
        "experimental_validation_used": False
    }, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "manifest": str(manifest),
                      "grades": dict(grades), "decisions": dict(decisions)}, indent=2))
if __name__ == "__main__": main()
